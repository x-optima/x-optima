#!/usr/bin/env python3
"""Build a 16:9 self-intro slide (PPTX + PNG) from the GitHub profile copy."""

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

OUT = Path(__file__).resolve().parent
W, H = 1920, 1080
BG = (15, 23, 42)
SURFACE = (30, 41, 59)
LINE = (51, 65, 85)
TEXT = (248, 250, 252)
MUTED = (148, 163, 184)
ACCENT = (45, 212, 191)
ACCENT_DIM = (17, 94, 89)
PITCH = (226, 232, 240)

FONT_DIR = Path("/usr/share/fonts/truetype/ubuntu")
FONT_BOLD = FONT_DIR / "Ubuntu-B.ttf"
FONT_REG = FONT_DIR / "Ubuntu-R.ttf"
FONT_MED = FONT_DIR / "Ubuntu-M.ttf"


def font(path: Path, size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(path), size)


def set_run(run, size, bold=False, color=TEXT, name="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    latin = rPr.find(qn("a:latin"))
    if latin is None:
        latin = rPr.makeelement(qn("a:latin"), {})
        rPr.append(latin)
    latin.set("typeface", name)


def add_textbox(slide, l, t, w, h, text, size, bold=False, color=TEXT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Emu(l), Emu(t), Emu(w), Emu(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color)
    return box


def add_rect(slide, l, t, w, h, fill, line=None, radius=None):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape, Emu(l), Emu(t), Emu(w), Emu(h))
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor(*fill)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = RGBColor(*line)
        s.line.width = Pt(1.25)
    if radius is not None:
        s.adjustments[0] = radius
    return s


def px_to_emu(px: int, slide_px: int, slide_inches: float) -> int:
    return int(px / slide_px * slide_inches * 914400)


def build_pptx() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*BG)

    def x(px):
        return px_to_emu(px, W, 13.333333)

    def y(px):
        return px_to_emu(px, H, 7.5)

    add_rect(slide, x(0), y(0), x(14), y(1080), ACCENT)

    add_textbox(slide, x(96), y(72), x(1100), y(90), "Виталий Кучин", 40, True)
    add_textbox(slide, x(96), y(158), x(800), y(48), "DevOps (практика)", 18, False, ACCENT)

    def pill(left, top, width, label):
        add_rect(slide, x(left), y(top), x(width), y(44), SURFACE, LINE, radius=0.5)
        add_textbox(slide, x(left), y(top + 6), x(width), y(32), label, 12, False, MUTED, PP_ALIGN.CENTER)

    pill(1320, 84, 160, "удалённо")
    pill(1496, 84, 220, "совместительство")
    pill(1732, 84, 120, "Россия")

    add_textbox(
        slide,
        x(96),
        y(248),
        x(1720),
        y(120),
        "С июля 2025 — практика DevOps: курс Нетологии и стенды в Yandex Cloud. "
        "Ищу удалённую команду: Terraform, Ansible, CI, мониторинг на виртуальных машинах.",
        18,
        False,
        PITCH,
    )

    tags = [
        (96, "Yandex Cloud", 180),
        (288, "Terraform", 150),
        (450, "Ansible", 130),
        (592, "GitLab CI", 150),
        (754, "Docker", 120),
        (886, "Prometheus", 160),
        (1058, "Grafana", 130),
        (1200, "PostgreSQL", 160),
    ]
    for left, label, width in tags:
        add_rect(slide, x(left), y(388), x(width), y(40), ACCENT_DIM, radius=0.2)
        add_textbox(slide, x(left), y(394), x(width), y(28), label, 12, True, ACCENT, PP_ALIGN.CENTER)

    cards = [
        (96, "ЛУЧШИЙ СТЕНД", "terraform6", "VPC, FastAPI в Docker, Managed MySQL, секреты в Lockbox, remote state."),
        (680, "CI В ОБЛАКЕ", "GitLab / TeamCity", "Self-hosted CI на ВМ. Maven: test вне main, deploy в Nexus с main."),
        (1264, "НАБЛЮДАЕМОСТЬ", "Prometheus + Grafana", "Метрики, дашборды, Alertmanager. Логи: ELK и Vector + ClickHouse."),
    ]
    for left, kicker, title, body in cards:
        add_rect(slide, x(left), y(468), x(560), y(430), SURFACE, LINE, radius=0.08)
        add_textbox(slide, x(left + 32), y(496), x(496), y(36), kicker, 11, True, ACCENT)
        add_textbox(slide, x(left + 32), y(536), x(496), y(56), title, 20, True)
        add_textbox(slide, x(left + 32), y(604), x(496), y(220), body, 14, False, MUTED)

    add_rect(slide, x(96), y(930), x(1728), y(2), LINE)
    add_textbox(
        slide,
        x(96),
        y(956),
        x(900),
        y(40),
        "github.com/x-optima   ·   Netology-DevOps",
        13,
        False,
        MUTED,
    )
    add_textbox(
        slide,
        x(900),
        y(956),
        x(924),
        y(40),
        "t.me/kuchinvn   ·   vkuchin-home@yandex.ru",
        13,
        False,
        MUTED,
        PP_ALIGN.RIGHT,
    )

    path = OUT / "vitaly-devops.pptx"
    prs.save(path)
    return path


def rounded_rect(draw, xy, radius, fill, outline=None, width=1):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def build_png() -> Path:
    img = Image.new("RGB", (W, H), BG)
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, 14, H), fill=ACCENT)

    name_f = font(FONT_BOLD, 64)
    role_f = font(FONT_MED, 28)
    pill_f = font(FONT_REG, 18)
    pitch_f = font(FONT_REG, 30)
    tag_f = font(FONT_MED, 18)
    kicker_f = font(FONT_MED, 15)
    card_title_f = font(FONT_BOLD, 30)
    card_body_f = font(FONT_REG, 20)
    foot_f = font(FONT_REG, 20)

    d.text((96, 72), "Виталий Кучин", font=name_f, fill=TEXT)
    d.text((96, 148), "DevOps (практика)", font=role_f, fill=ACCENT)

    pills = [("удалённо", 160), ("совместительство", 220), ("Россия", 120)]
    x = 1732 + 120
    for label, w in reversed(pills):
        x -= w
        rounded_rect(d, (x, 84, x + w, 128), 22, SURFACE, LINE, 2)
        bbox = d.textbbox((0, 0), label, font=pill_f)
        tw = bbox[2] - bbox[0]
        d.text((x + (w - tw) / 2, 94), label, font=pill_f, fill=MUTED)
        x -= 16

    pitch = (
        "С июля 2025 — практика DevOps: курс Нетологии и стенды в Yandex Cloud.\n"
        "Ищу удалённую команду: Terraform, Ansible, CI, мониторинг на виртуальных машинах."
    )
    d.text((96, 248), pitch, font=pitch_f, fill=PITCH, spacing=12)

    tags = ["Yandex Cloud", "Terraform", "Ansible", "GitLab CI", "Docker", "Prometheus", "Grafana", "PostgreSQL"]
    tx = 96
    ty = 388
    for label in tags:
        bbox = d.textbbox((0, 0), label, font=tag_f)
        tw = bbox[2] - bbox[0] + 28
        th = 40
        rounded_rect(d, (tx, ty, tx + tw, ty + th), 8, ACCENT_DIM)
        d.text((tx + 14, ty + 8), label, font=tag_f, fill=ACCENT)
        tx += tw + 12

    cards = [
        (96, "ЛУЧШИЙ СТЕНД", "terraform6", "VPC, FastAPI в Docker, Managed MySQL,\nсекреты в Lockbox, remote state."),
        (680, "CI В ОБЛАКЕ", "GitLab / TeamCity", "Self-hosted CI на ВМ. Maven: test вне main,\ndeploy в Nexus с main."),
        (1264, "НАБЛЮДАЕМОСТЬ", "Prometheus + Grafana", "Метрики, дашборды, Alertmanager.\nЛоги: ELK и Vector + ClickHouse."),
    ]
    for left, kicker, title, body in cards:
        rounded_rect(d, (left, 468, left + 560, 898), 16, SURFACE, LINE, 2)
        d.text((left + 32, 496), kicker, font=kicker_f, fill=ACCENT)
        d.text((left + 32, 536), title, font=card_title_f, fill=TEXT)
        d.text((left + 32, 604), body, font=card_body_f, fill=MUTED, spacing=8)

    d.rectangle((96, 930, 1824, 932), fill=LINE)
    d.text((96, 956), "github.com/x-optima   ·   Netology-DevOps", font=foot_f, fill=MUTED)
    right = "t.me/kuchinvn   ·   vkuchin-home@yandex.ru"
    rb = d.textbbox((0, 0), right, font=foot_f)
    d.text((1824 - (rb[2] - rb[0]), 956), right, font=foot_f, fill=MUTED)

    path = OUT / "vitaly-devops.png"
    img.save(path, "PNG")
    return path


if __name__ == "__main__":
    pptx = build_pptx()
    png = build_png()
    print(pptx)
    print(png)
